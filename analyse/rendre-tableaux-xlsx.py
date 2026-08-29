"""Rassemble les tableaux de la partie III dans un seul classeur Excel.

    "C:/Users/videt/anaconda3/python.exe" analyse/rendre-tableaux-xlsx.py

Lit `memoire/tableaux/*.csv` (produits par `analyse/R/analyse.R`) et écrit
`memoire/tableaux-partie3.xlsx` : un onglet par tableau, en-têtes en gras, volets
figés, colonnes dimensionnées, nombres alignés à droite.

Le CSV reste la source. Une retouche faite dans le classeur sera écrasée à la
prochaine régénération : la reporter dans le script R.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).resolve().parent.parent
TABLEAUX = RACINE / "memoire" / "tableaux"
SORTIE = RACINE / "memoire" / "tableaux-partie3.xlsx"

#: Ordre de lecture du mémoire, et non ordre alphabétique des fichiers.
ORDRE = [
    "T1-serie-decomposee",
    "T2-echelle-de-lecture",
    "T3-effet-apparie-H1",
    "T4-reference-recitee",
    "T5-effet-apparie-H3",
    "T6-quatre-decisions-Station",
    "T7-vocabulaire-notes",
    "T7bis-registre-par-adversaire",
    "T8-degenerescence",
    "T9-incoherence",
    "T10-synthese-hypotheses",
]

GRIS = PatternFill("solid", fgColor="EDEDED")
TRAIT = Side(style="thin", color="BFBFBF")
BORDURE = Border(left=TRAIT, right=TRAIT, top=TRAIT, bottom=TRAIT)

#: Au-delà, la colonne passe en retour à la ligne plutôt que de s'étirer.
LARGEUR_MAX = 60


def nombre(texte: str):
    """Rend un float si la cellule en est un, sinon le texte tel quel."""
    try:
        return float(texte)
    except (TypeError, ValueError):
        return texte


def ajouter(classeur: Workbook, chemin: Path) -> int:
    # utf-8-sig : les CSV sont écrits avec BOM pour qu'Excel les ouvre droit.
    with chemin.open(encoding="utf-8-sig", newline="") as fh:
        lignes = list(csv.reader(fh))
    if not lignes:
        return 0

    # 31 caractères est la limite d'Excel pour un nom d'onglet.
    feuille = classeur.create_sheet(chemin.stem[:31])
    for rangee in lignes:
        feuille.append([nombre(c) for c in rangee])

    for cellule in feuille[1]:
        cellule.font = Font(bold=True)
        cellule.fill = GRIS
        cellule.alignment = Alignment(vertical="center", wrap_text=True)

    for rangee in feuille.iter_rows():
        for cellule in rangee:
            cellule.border = BORDURE
            if isinstance(cellule.value, float):
                cellule.number_format = "0.0000"
                cellule.alignment = Alignment(horizontal="right")
            elif cellule.row > 1:
                cellule.alignment = Alignment(vertical="top", wrap_text=True)

    for i, colonne in enumerate(feuille.columns, start=1):
        large = max(len(str(c.value)) if c.value is not None else 0 for c in colonne)
        feuille.column_dimensions[get_column_letter(i)].width = min(large + 3, LARGEUR_MAX)

    feuille.freeze_panes = "A2"
    return len(lignes) - 1


def main() -> None:
    classeur = Workbook()
    classeur.remove(classeur.active)

    presents = {p.stem: p for p in sorted(TABLEAUX.glob("*.csv"))}
    manquants = [n for n in ORDRE if n not in presents]
    orphelins = [n for n in presents if n not in ORDRE]

    total = 0
    for nom in ORDRE:
        if nom in presents:
            n = ajouter(classeur, presents[nom])
            total += n
            print(f"  {nom:32s} {n:3d} lignes")
    for nom in orphelins:  # un tableau ajouté au script R et pas encore listé
        n = ajouter(classeur, presents[nom])
        total += n
        print(f"  {nom:32s} {n:3d} lignes  (hors ordre de lecture)")

    if manquants:
        print("\n  ATTENTION — tableaux attendus et absents :")
        for nom in manquants:
            print(f"    {nom}  — relancer analyse/R/analyse.R")

    classeur.save(SORTIE)
    print(f"\nécrit : {SORTIE.relative_to(RACINE)}")
    print(f"  {len(classeur.sheetnames)} onglets · {total} lignes de données")


if __name__ == "__main__":
    main()
