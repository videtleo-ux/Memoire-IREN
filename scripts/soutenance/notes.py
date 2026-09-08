# -*- coding: utf-8 -*-
"""Notes de présentateur, une par slide. Texte suivi, à dire tel quel. Cible : 20 minutes."""
import re
from pathlib import Path

from pptx import Presentation
from pptx.oxml.ns import qn

PPTX = Path(__file__).resolve().parent.parent.parent / "Soutenance.pptx"

NOTES = {
1: """[~30 s]

Bonjour. Je vais vous présenter un travail qui pose une question simple : quand on donne
une mémoire à un agent doté d'un modèle de langage, et qu'on le place dans une situation
stratégique où il faut deviner ce que fait l'autre, développe-t-il quelque chose qui
mérite le nom de stratégie ?

Je procéderai en trois temps : d'où vient la question, comment je l'ai rendue mesurable,
et ce que les mesures disent. J'ai construit une arène expérimentale pour cela, et j'y ai
fait jouer un agent sur plus de vingt-sept mille décisions.""",

2: """[~1 min 40]

Les modèles de langage ne produisent pas seulement du texte. Ils sont déployés comme
agents autonomes : ils négocient des contrats, enchérissent, fixent des prix, arbitrent
des allocations, et de plus en plus souvent face à d'autres agents dont les intérêts
divergent des leurs. L'enjeu est économique avant d'être technique : si ces agents
prennent une part croissante des décisions de marché, alors la nature de leur rationalité
détermine les équilibres qui s'y forment.

Ce n'est pas une projection. Je l'ai observé pendant deux ans en conseil en innovation, et
toujours dans le même ordre.

D'abord l'IA outille la décision. Chez FDJ United, j'ai contribué à un dispositif de veille
où l'analyse automatique filtre plus de cent articles par semaine avant que nous les
éditorialisions.

Ensuite elle entre dans les process. Dans une maison de luxe, j'ai conduit quarante-quatre
entretiens, de la vente au comité exécutif, pour aboutir à cinq cas d'usage prototypés et
chiffrés pour le retail. Dans un groupe industriel, la question posée au comité exécutif
n'était déjà plus l'outil, mais la feuille de route.

Puis elle agit seule. J'ai construit un assistant qui cartographie une chaîne de valeur
sans moi : cinq minutes, contre plusieurs heures. Il ne me répond pas, il produit.

C'est à ce dernier palier que ma question se pose, et personne autour de la table ne savait
y répondre : sur quoi ces agents fondent-ils leurs décisions ? D'où l'énoncé du bandeau. Il
a une seconde branche, que je garde pour la fin : cette adaptation rapproche-t-elle l'agent
de l'optimalité, ou reproduit-elle des biais ?""",

3: """[~1 min 50]

Trois littératures, et elles ne jouent pas le même rôle : la première fournit la méthode,
la deuxième le sujet, la troisième le terrain.

L'économie comportementale et expérimentale substitue à l'Homo œconomicus un agent
psychologiquement situé, comme le posent Jacquemet et L'Haridon. Elle m'apporte surtout un
instrument : le laboratoire comme moyen de preuve. Plutôt que de décrire un comportement,
on le confronte à une norme théorique et on mesure l'écart, si bien que le biais devient un
objet structuré, un écart systématique à une norme, et non une erreur aléatoire. Cot et
Ferey ajoutent un avertissement que je reprends en conclusion, avec leur scission
tripartite de la rationalité : l'optimalité n'est pas univoque, elle dépend du cadre qui la
juge.

Le deuxième pôle transpose ce laboratoire aux machines. Rahwan et ses coauteurs proposent
d'étudier les machines comme une classe d'acteurs, et non comme du code : leur comportement
émergent n'est pas formalisable analytiquement, il faut donc l'expérimenter. Horton et ses
coauteurs vont plus loin, en traitant le modèle de langage comme un modèle computationnel
implicite de l'humain, doté par le langage d'une dotation et de croyances. C'est là que je
me démarque : on peut simuler l'humain, ou observer une entité artificielle pour elle-même.
Je fais le second choix, et c'est pourquoi je ne conclurai pas sur ce que l'agent partage
avec nous.

Le terrain, enfin, est le poker de Kuhn, et il vient d'une filiation directe : Borel, von
Neumann, puis Kuhn. Ce qu'elle établit tient en deux points. Le bluff n'est pas une ruse
psychologique, c'est une conduite rationnelle. Et l'équilibre y est en stratégies mixtes.
Retenez ce second point : c'est lui qui reviendra à ma dernière hypothèse.""",

4: """[~1 min 40]

Il existe trois façons d'adapter un modèle sans le réentraîner, et chacune a son défaut.
La réinjection dans la fenêtre de contexte en fait une mémoire de travail : c'est flexible,
mais volatile, tout s'effondre à la fermeture de la session. La gestion sémantique résume
pour ne pas saturer, au prix des micro-signaux quantitatifs, comme le montrent Chen et ses
coauteurs — or ce sont exactement ceux dont on a besoin en stratégie fine. L'apprentissage
par renforcement et le réglage fin sont permanents, mais rigides, coûteux, et ils imitent
passivement les biais du jeu de données.

La mémoire auto-rédigée lève la contradiction. Voyager stocke l'expérience en programmes
réutilisables, jamais en poids. Reflexion fait analyser son échec à l'agent, lui fait
rédiger une auto-réflexion en langue naturelle, qu'il relit au tour suivant. L'agent écrit
donc lui-même ce qu'il retient, sans qu'on touche à ses paramètres.

Voici la charnière de mon travail : ces architectures ont toutes été validées dans des
environnements qui ne cherchent jamais à tromper l'agent. Minecraft ne bluffe pas.

Sur le terrain adversarial, les travaux existants ratent chacun une pièce. Richelieu
travaille sur une information quasi publique. MemGPT est bien évalué sur le poker de Kuhn,
mais avec des critères d'ingénierie logicielle. Readable Minds porte sur un jeu où
l'équilibre n'est pas calculable, donc sans étalon. Et tous confondent trois dimensions que
je vais instrumenter séparément : le biais de décision, la stratégie d'équilibre et la
stratégie exploitative.

D'où le positionnement, en une ligne. Pour exploiter, il faut estimer la déviation de
l'adversaire ; pour l'estimer, il faut comparer des manches entre elles ; et pour les
comparer, il faut une mémoire. Sans mémoire persistante, un agent est condamné à
l'équilibre défensif.""",

5: """[~2 min]

La tâche est le poker de Kuhn : trois cartes, un jeton d'ante, une seule mise possible,
douze ensembles d'information. C'est assez petit pour que l'équilibre se calcule
exactement, et assez riche pour que le bluff y soit optimal.

Le vocabulaire est réécrit partout. Aucune mention de poker, de carte, de roi ou de valet :
des sceaux qui s'appellent Tor, Vael et Rhun, des actions retenir, engager, couvrir, se
retirer. La raison est que le modèle a rencontré ce jeu et sa solution pendant son
entraînement. Je ne peux pas garantir qu'il ne le reconnaît pas, et je le mesure
explicitement, mais je ne le lui sers pas sur un plateau.

Les trois adversaires sont le cœur du plan expérimental. GTO joue l'équilibre exact, à
alpha égale un tiers : il n'y a rien à y exploiter, c'est mon contrôle négatif. Station ne
mise jamais et paie toujours, il faut donc cesser de bluffer contre elle. Over-folder ne
mise jamais et se couche toujours, il faut donc engager toutes les mains. Ces deux fuites
sont exactement opposées : aucune stratégie fixe, l'équilibre compris, ne peut faire
baisser les deux écarts ensemble. C'est ce qui sépare exploiter de réciter.

Les trois traitements, ensuite. Le modèle est tenu constant, le harnais aussi, le texte des
règles aussi : seule la rubrique mémoire du prompt change. Sans mémoire, elle est vide. En
réinjection, l'expérimentateur empile les récapitulatifs passés dans une fenêtre de
cinquante-deux mille caractères, dimensionnée pour contenir trois séries — donc pour que
cette condition puisse accumuler, et aussi saturer. En auto-écriture, l'agent lit le
récapitulatif et réécrit lui-même son fichier de notes.

Une série, enfin, c'est cent cinquante manches jouées sous un état mémoire fixe, pour qu'un
point de courbe corresponde à un état mémoire et à un seul. La mémoire est capturée à
l'ouverture, figée pendant toute la série, mise à jour à la seule frontière, avec un
contexte neuf à chaque manche. Et j'insiste sur la dernière ligne : les écritures
intra-série sont détectées, enregistrées, puis annulées — le gel est vérifié, pas supposé.
Enfin, les trois traitements d'une même réplication reçoivent les mêmes distributions,
manche par manche. Les comparaisons se font à chance égale.""",

6: """[~1 min 50]

L'écart d'exploitation est la différence entre ce que rapporterait la meilleure réponse à
cet adversaire et ce que rapporte la politique observée. Ce n'est ni le gain ni le taux de
victoire, deux quantités qui confondent la qualité du jeu et la chance de la distribution.
Il est exact — six distributions, un arbre de profondeur trois, énumération complète —
donc il n'y a ni simulation ni intervalle de confiance sur la mesure elle-même. Il est
insensible à la chance : les résultats portent sur la politique, pas sur les gains
réalisés. Et il se lit en jetons par manche, zéro signifiant exploitation optimale.

À côté, je calcule la même espérance contre une autre référence, l'équilibre. Un récitant y
reste à zéro quel que soit l'adversaire, tandis qu'un exploiteur s'en écarte positivement :
c'est l'instrument de ma deuxième hypothèse. Par honnêteté, je précise qu'à adversaire fixé
les deux mesures n'ont qu'un seul degré de liberté. Ce n'est pas une seconde mesure
indépendante, c'est un recalage : il déplace le zéro sur le comportement du récitant, ce
qui rend lisible le franchissement de signe.

Le tableau donne l'échelle de lecture, parce qu'un écart ne s'interprète pas dans l'absolu :
le plafond dépend de l'adversaire. Contre Over-folder, un joueur parfait gagne un jeton par
manche là où un récitant n'en gagne que 0,22 — la référence récitée peut donc monter à plus
de 0,78. Contre Station, elle plafonne à 0,11. Contre GTO, elle vaut zéro, et c'est
exactement pour cela que GTO sert de contrôle.

La campagne, enfin : vingt-quatre exécutions, cent soixante-dix-sept séries, 26 550 manches,
27 290 décisions de l'agent, sur gpt-5.6-luna à effort medium, cent cinquante manches par
série, trois réplications, une graine de campagne unique. Quarante-sept dollars, deux jours,
une seule machine. Et la vérification : deux cent vingt-cinq tests automatisés, dont un
oracle analytique du moteur de mesure, plus un rejeu qui re-règle chaque manche depuis les
seuls journaux et retrouve les résultats publiés, vingt-quatre fois sur vingt-quatre.
L'expérience est reconstructible.""",

7: """[~1 min 20]

Quatre hypothèses. Je ne les lis pas, j'en donne la logique : chacune écarte une
explication concurrente, et chacune est réfutable.

La première pose l'existence de l'effet et son origine. L'écart doit décroître pour l'agent
à mémoire auto-rédigée et rester stationnaire pour le même modèle privé de mémoire. C'est
le témoin sans mémoire qui fait le travail d'identification : mêmes cartes, même modèle,
même calcul, seule la rubrique mémoire change. Elle tombe si les deux conditions ne se
distinguent pas.

La deuxième est la plus importante, parce qu'elle répond à l'objection sérieuse. On peut
m'objecter que l'effet est trivial : le modèle connaît la solution du jeu, et la mémoire
l'aide simplement à s'en souvenir. Ce serait une récitation mieux exécutée, pas une
adaptation. L'hypothèse exige donc que l'agent s'écarte de l'équilibre dans la direction
que commande l'adversaire rencontré — cesser de bluffer contre l'un, engager toutes les
mains contre l'autre. Elle tombe si l'écart ne descend que contre un seul des deux
adversaires biaisés.

La troisième demande si le mécanisme de rétention compte, à information strictement
identique : la réinjection reçoit exactement la même matière que l'auto-écriture, le
récapitulatif brut. Ce qui diffère, c'est l'étape de synthèse.

La quatrième est celle qui pouvait retourner tout le reste. Elle demande si l'agent
conserve des défaillances alors même que sa politique s'améliore, et je la mesure d'abord
sur la cohérence entre ce qu'il raisonne et ce qu'il joue.""",

8: """[~1 min 20]

La figure porte l'écart série par série, un panneau par adversaire, les trois traitements
sur les mêmes axes. La courbe sans mémoire s'arrête à trois points, parce que le protocole
n'en prévoit que trois : rien ne s'y accumule.

Un seul point à faire ici, et il se voit. La ligne sans mémoire ne décroît jamais, contre
aucun des trois adversaires ; sa pente est même légèrement positive contre deux d'entre
eux. La ligne à mémoire auto-écrite, elle, chute dès la première frontière et atteint zéro
contre les deux adversaires exploitables.

Les trois tuiles donnent l'effet apparié, à cartes identiques. Contre Over-folder, l'écart
passe de 0,704 à zéro, et les trois réplications sur trois vont dans le même sens. Contre
Station, de 0,222 à 0,002, avec deux réplications à zéro exact. Contre GTO, de 0,187 à
0,008.

Je garde deux précisions pour les questions. D'abord, zéro veut dire zéro exactement, et
non proche de zéro : la mesure étant exacte, un écart nul signifie que la politique
observée est une meilleure réponse à l'adversaire, au sens strict. Ensuite, si cette
décroissance venait de l'instrument, de la séquence de cartes ou d'un artefact du harnais,
elle apparaîtrait aussi dans la condition sans mémoire. Elle n'y apparaît pas.""",

9: """[~1 min 50]

L'objection est la suivante : le modèle connaît peut-être la solution du jeu, et la mémoire
l'aide simplement à mieux la réciter. La figure tranche sans passer par aucune mesure de
gain. Une seule situation y est représentée : l'agent détient le sceau le plus faible, et
on regarde à quelle fréquence il l'engage.

Sans mémoire, cette fréquence ne peut pas dépendre de l'adversaire, et c'est une propriété
du dispositif avant d'être une observation : en position d'ouvrant, l'agent parle le
premier, aucune action adverse n'a eu lieu, le contexte est neuf à chaque manche, donc son
information est rigoureusement identique face aux trois. C'est bien ce qu'on observe, avec
0,285, 0,307 et 0,355, pour un écart-type d'échantillonnage de 0,031. Rien ne les distingue.

Avec mémoire, la même fréquence monte à 1,000 contre Over-folder et tombe à 0,004 contre
Station. Trois réponses à la même situation, dont deux diamétralement opposées.

Aucune application de l'équilibre ne produit ce contraste. L'équilibre est un objet fixe,
douze probabilités calculées contre un adversaire supposé jouer lui-même l'équilibre. Un
agent qui révise ses croyances comme le prescrit l'équilibre bayésien parfait n'en sort pas,
parce que ces croyances portent sur le nœud atteint, pas sur la stratégie de l'adversaire.
Réviser ses croyances sur la stratégie adverse suppose de comparer des manches entre elles,
donc une mémoire, et c'est exactement ce que j'appelle exploiter.

Le tableau donne l'amplitude. Contre les deux adversaires exploitables, l'agent à mémoire
atteint 0,778 et 0,110 pour des maxima théoriques de 0,778 et 0,111 : cent pour cent et
quatre-vingt-dix-neuf pour cent de l'exploitation disponible. Et je vous fais remarquer la
colonne sans mémoire : à moins 0,111 contre Station et moins 0,187 contre GTO, il est en
dessous du niveau de l'équilibre.

Le contrôle négatif, pour finir. Contre l'adversaire à l'équilibre, aucune politique ne peut
rapporter plus que l'équilibre. Sur trente-neuf séries, la référence récitée n'est jamais
positive et l'écart jamais négatif ; six séries atteignent exactement zéro, aucune ne va
au-delà. L'instrument tient.""",

10: """[~2 min]

À information égale, deux mécanismes de rétention produisent-ils la même adaptation ? La
réponse est oui pour l'essentiel, et non sur un point précis — et ce point n'est ni celui
que j'anticipais, ni là où je le cherchais.

Deux prédictions ne se réalisent pas. J'annonçais pour la réinjection un profil d'oubli en
dents de scie quand la fenêtre sature : il n'y en a aucune trace. Et je m'attendais à une
différence de vitesse : il n'y en a pas non plus, les deux conditions chutent à la première
frontière dans les mêmes proportions.

Ce qui les sépare est ailleurs, et c'est le tableau qui le montre, parce que l'écart est une
quantité agrégée qui masque où se loge le résidu. Quatre décisions décident de tout contre
Station. Sur les deux premières, la meilleure réponse est de ne jamais engager le sceau
faible : la réinjection le fait encore dans six et neuf pour cent des cas, quand
l'auto-écriture est retombée sous le pour cent. Sur les deux suivantes, il faut engager le
sceau fort systématiquement : la réinjection laisse passer entre trois et six pour cent de
ces engagements de valeur, quand l'auto-écriture est à 1,000, exactement.

Sur l'écart agrégé, la différence entre les deux vaut 0,038 contre Station, avec un
intervalle de confiance à quatre-vingt-quinze pour cent de plus ou moins 0,015 qui exclut
zéro. Contre Over-folder, elle vaut 0,004, et l'intervalle contient zéro.

D'où une conclusion que j'énonce moi-même comme conditionnelle : ma troisième hypothèse est
vérifiée, mais le mécanisme ne départage les deux mémoires que contre l'un des deux
adversaires exploitables.

Et c'est précisément ce qui l'explique. Contre Over-folder, la politique optimale tient en
une seule règle, positive : engager, toujours. Les deux mécanismes l'extraient intégralement.
Contre Station, il en faut deux, dont une négative — et cesser de faire quelque chose est ce
qu'un historique brut soutient le plus mal : un engagement perdant y figure comme une manche
parmi cent cinquante, noyée dans celles où le même geste a gagné, et rien ne rassemble ces
occurrences en interdiction. L'étape de synthèse fait exactement cela : elle contraint
l'agent à écrire une règle, et une règle énonce ce qu'il ne faut pas faire.

L'apport de la mémoire auto-écrite est donc plus étroit que mon hypothèse ne le supposait :
elle ne retient ni mieux ni plus longtemps, elle achève l'induction.""",

11: """[~1 min 50]

Jusqu'ici, un agent qui s'améliore. Ma quatrième hypothèse pose la question inverse, et
c'est là que les deux branches de la problématique se rejoignent.

Contre un adversaire à l'équilibre, bien jouer exige d'engager le sceau faible une fois sur
trois, et de façon imprévisible. Toujours l'engager ou ne jamais l'engager se punit
également. C'est le point que j'annonçais en revue de littérature : l'équilibre de ce jeu
est en stratégies mixtes, donc l'imprévisibilité y est la condition de l'optimalité.

Sur la figure, chaque point est une série. Sans mémoire, les fréquences se tiennent autour
de vingt-huit pour cent, pour un optimum à trente-trois. Dès qu'une note existe, elles se
dispersent et se portent aux valeurs extrêmes plus d'une série sur deux. L'agent n'a pas
oublié comment jouer, puisque son écart continue de baisser : il a cessé de tirer au sort.

Le tableau donne le mécanisme, et il vient d'une lecture intégrale des quatre-vingt-dix
notes, jamais éditées. Aucun pourcentage. Aucune proportion en toutes lettres. Aucune
fréquence prescrite pour sa propre action. En regard, deux cent vingt et un quantificateurs
catégoriques — toujours, jamais — et deux cent vingt-huit quantificateurs gradués. Il n'est
pas incapable d'écrire une proportion, il le fait pour décrire une probabilité d'abattage.
Il ne le fait jamais pour prescrire sa propre fréquence. Ce n'est pas le vocabulaire qui
manque, c'est l'usage.

D'où ma conclusion : c'est le même mécanisme. Ce qui rend la mémoire auto-écrite supérieure,
c'est-à-dire la contrainte de formuler une règle, est exactement ce qui la rend incapable de
randomiser, parce qu'une règle écrite en langue naturelle est déterministe. L'adaptation et
le biais ne coexistent pas par accident.

J'énonce la réserve moi-même : le coût est déduit, pas mesuré. Face aux bots figés de la
campagne, abandonner le mélange ne coûte rien, c'est un point d'indifférence exact. Ce qu'il
coûterait face à un adversaire capable de s'ajuster relève de la théorie, et cet adversaire
est hors périmètre.

La formule à retenir est donc qu'un agent doté de mémoire est un meilleur exploiteur et un
adversaire plus prévisible. Dans une interaction répétée avec un partenaire qui apprend, la
seconde propriété peut annuler la première.""",

12: """[~1 min 20]

Je préfère énoncer les limites moi-même, franchement, plutôt que de les laisser venir en
question.

Le dispositif est restreint : trois réplications, un modèle, une tâche, un horizon de dix
séries. L'appariement compense la faiblesse de l'effectif, il ne l'annule pas. Le plateau
étant atteint dès les premières frontières, je mesure bien la persistance de l'adaptation
mais je résous mal sa vitesse : les séries suivantes documentent un état stable, pas la
dynamique qui y conduit. Le contrôle négatif n'a pas été conduit en condition de
réinjection, arbitrage de dimensionnement que j'assume. Et la dégénérescence repose sur un
seul ensemble d'information.

J'ajoute une limite qui ne figure pas au tableau, parce qu'elle porte sur la portée de mes
conclusions : il n'y a pas de bras humain. Le second terme de ma problématique, celui qui
demandait si l'adaptation reproduit les biais humains, n'a donc pas de comparateur mesuré,
et je ne qualifierai pas de biais humain un comportement qu'aucun humain n'a produit ici.
Ce que j'affirme est plus étroit que la question, et plus solide : l'agent construit une
stratégie, il l'optimise correctement, et il le fait dans un espace de stratégies pures
parce que le canal qui porte son adaptation ne sait rien écrire d'autre.

Les prolongements suivent ces limites une à une. Des séries plus courtes et plus nombreuses
pour cerner la vitesse et l'oubli. Un jeu comportant plusieurs ensembles d'information à
équilibre mixte, et plusieurs modèles. Un adversaire biaisé puis à l'équilibre, mémoire
conservée, pour savoir si l'exploitation apprise devient une charge. Le même protocole joué
par des humains, pour la référence manquante. Et l'adversaire lui-même doté d'un modèle de
langage, ce qui ferait de la prévisibilité une variable dynamique plutôt qu'une observation.

Merci. Je suis à votre disposition.""",
}


def sans_puce(paragraphe) -> None:
    """Le notesMaster du squelette impose une puce ● à chaque paragraphe.

    Les notes sont de la prose : on la supprime en posant <a:buNone/> sur le
    paragraphe, qui l'emporte sur l'héritage. buNone doit suivre les attributs
    d'indentation dans pPr, d'où l'insertion en tête (pPr n'en porte pas ici).
    """
    pPr = paragraphe._p.get_or_add_pPr()
    for balise in ("buChar", "buAutoNum", "buNone"):
        for vieux in pPr.findall(qn(f"a:{balise}")):
            pPr.remove(vieux)
    pPr.insert(0, pPr.makeelement(qn("a:buNone"), {}))


def deplier(texte: str) -> list[str]:
    """Recolle les paragraphes du source, coupé à ~90 colonnes pour la lisibilité.

    Sans cela, chaque ligne du source devient un paragraphe PowerPoint — donc une
    puce et un retour forcé au milieu d'une phrase. Un paragraphe = un bloc séparé
    par une ligne vide.
    """
    blocs = re.split(r"\n\s*\n", texte.strip())
    return [" ".join(bloc.split()) for bloc in blocs if bloc.strip()]


def main() -> None:
    pres = Presentation(PPTX)
    for numero, texte in NOTES.items():
        cadre = pres.slides[numero - 1].notes_slide.notes_text_frame
        cadre.text = "\n".join(deplier(texte))
        for paragraphe in cadre.paragraphs:
            sans_puce(paragraphe)
    pres.save(PPTX)
    print("notes écrites :", len(NOTES), "slides")


if __name__ == "__main__":
    main()
